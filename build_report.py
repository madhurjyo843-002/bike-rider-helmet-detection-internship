from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

base = Path(r'D:\internship project')
out = base / 'Internship_Project_Report.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.72)
sec.left_margin = Inches(.82); sec.right_margin = Inches(.82)
styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Aptos'; normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(45,55,67)
normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.05
for name, size, color in [('Title',28,'16324F'),('Heading 1',17,'16324F'),('Heading 2',12,'087E8B'),('Heading 3',10.5,'16324F')]:
    s=styles[name]; s.font.name='Aptos Display' if name in ('Title','Heading 1') else 'Aptos'
    s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(5)
title_ppr=styles['Title']._element.get_or_add_pPr()
title_border=title_ppr.find(qn('w:pBdr'))
if title_border is not None: title_ppr.remove(title_border)
header=sec.header.paragraphs[0]; header.text='INTERNSHIP PROJECT REPORT  |  COMPUTER VISION'
header.style='Caption'; header.runs[0].font.color.rgb=RGBColor(95,112,128); header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Motorcycle Rider Helmet Detection  •  '); r.font.size=Pt(8); r.font.color.rgb=RGBColor(95,112,128)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def para(text='', style=None):
    p=doc.add_paragraph(style=style); p.add_run(text); return p
def h(text, level=1): doc.add_heading(text, level=level)
def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent=Inches(.25)
    p.paragraph_format.space_after=Pt(3); p.add_run(text); return p
def shade(cell, fill):
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); cell._tc.get_or_add_tcPr().append(shd)
def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,x in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=x; shade(c,'16324F')
        for p in c.paragraphs: p.paragraph_format.keep_with_next=True
        for rr in c.paragraphs[0].runs: rr.bold=True; rr.font.size=Pt(9); rr.font.color.rgb=RGBColor(255,255,255)
    header_props=t.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); header_props.append(repeat)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,x in enumerate(row):
            cells[i].text=str(x); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
                for rr in p.runs: rr.font.name='Aptos'; rr.font.size=Pt(9)
            if ri%2: shade(cells[i],'F1F6F8')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    for row in t.rows:
        trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
    spacer=doc.add_paragraph(); spacer.paragraph_format.space_after=Pt(0); spacer.paragraph_format.line_spacing=Pt(1)
    spacer.add_run(' ').font.size=Pt(1)
    return t

# Cover
para('INTERNSHIP REPORT', 'Subtitle')
p=doc.add_paragraph(style='Title'); p.paragraph_format.space_before=Pt(52); p.add_run('Motorcycle Rider Helmet Detection')
para('A computer vision dataset assessment and research development plan','Subtitle')
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(25)
r=p.add_run('Submitted by'); r.bold=True; r.font.color.rgb=RGBColor(8,126,139)
for label,value in [('Student name','[Enter your name]'),('Student ID','[Enter student ID]'),('Program and department','[Enter program and department]'),('Host organization','[Enter company or internship host]'),('Department / supervisor','[Enter department and supervisor]'),('Internship period','[Enter start and end dates]'),('Submission date','25 October 2026')]:
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(4); r=p.add_run(label+': '); r.bold=True; p.add_run(value)
para('Project focus: preparation and critical assessment of a consented YOLO-format dataset for helmet and no-helmet object detection.','Caption')
doc.add_page_break()

h('Executive Summary')
para('This report documents a practical research-development project using a supplied bike-rider helmet detection dataset. The work completed for this report was a reproducible structural audit of the extracted dataset, a review of its label format and provenance, and a proposed experiment plan for YOLOv8 object detection. The audit counted 903 images in the supplied train, validation, and test splits. It found 1,701 bounding-box rows and one polygon annotation, no missing image-label pairs, no unreadable images, and no exact duplicate image files across splits.')
para('The archive does not contain trained model weights, and the available Python environment does not include PyTorch or Ultralytics. As a result, this work did not train or evaluate a model. No precision, recall, mAP, or detection claims are made. The most useful next steps are to review the mixed polygon/box annotation, define exactly what each class box represents, check source-level leakage and label consistency, then train a small baseline and report held-out test metrics.')
para('The report follows the requested internship structure. Company-specific facts and the student’s identifying information were not supplied, so bracketed fields are provided rather than invented. Replace these fields with verified information before submission.')

h('1. Introduction')
para('Road safety is a practical setting for computer vision because cameras can observe many road users at once, while a human reviewer may have limited time. This project examines one narrow task: detecting objects labelled as helmet or no_helmet in images of riders. A detector could help flag scenes for review, but an image model cannot by itself establish who a person is, whether a legal rule applies, or whether a rider is compliant in every real-world situation.')
para('The project objective is to assess whether the supplied dataset is organized and annotated well enough to support a defensible YOLOv8 detection experiment, identify risks that could distort results, and define how such an experiment should be carried out. The dataset was supplied by a friend who consented to its use for this project. The archive itself attributes the data to a Roboflow Universe project and marks it CC BY 4.0; attribution and license conditions should be preserved when sharing outputs.')
para('The scope completed here is data inspection and research planning. It includes image and label pairing, basic label-format checks, image-size counts, class counts, and exact-duplicate checks. It does not include model training, prediction, external validation, or a deployment assessment.')
para('The guiding research question is: can a YOLOv8 detector learn to distinguish the two supplied labels, and what additional evidence would be needed before trusting it on unseen road scenes? The present project answers the preparation part of that question. It establishes a baseline dataset profile and identifies checks that should be completed before model results are interpreted.')

h('2. Company Profile')
para('Host organization: [Enter verified organization name]. This project report does not assume a company or institutional host because no organization profile was included with the dataset. Add the host’s approved description here, including its sector, location, services, and relevance to the internship.')
para('Suggested project context to adapt: the work sits at the intersection of applied computer vision, data preparation, and software prototyping. If the internship host operates in transport, public safety, AI, or software development, explain that connection using the host’s own verified description. Do not present this suggested wording as a factual company profile until it is checked.')

h('3. Organizational Structure')
p=para('Insert the organization’s actual structure and reporting line here. A suitable description should name the intern’s department, supervisor role, technical team, and any review or approval roles involved in the project. No reporting relationship was provided, so none is claimed in this report.')
p.paragraph_format.keep_with_next=True
table(['Role','Name or title to complete','Contribution to confirm'],[
('Intern / student','[Your name and program]','Dataset audit, documentation, and proposed experiment'),
('Industry supervisor','[Name and position]','[Mentoring, review, or project direction]'),
('Academic supervisor','[Name and position]','[Academic guidance and assessment]'),
('Host team','[Department or unit]','[Data, infrastructure, or domain input]')],[1.35,2.05,3.55])

h('4. Technologies Used')
para('The dataset is organized for YOLO-style object detection: each image has a text label file containing a class identifier and normalized coordinates. The archive’s data.yaml lists two classes in this order: class 0 is helmet and class 1 is no_helmet. The images are 640 by 640 pixels. The accompanying Roboflow notes state that the export was generated on 16 May 2026, with auto-orientation and resize-to-640-by-640 stretching and no image augmentation.')
para('For a standard detection row, the label stores a class ID followed by the box centre coordinates, width, and height, normalized to the image dimensions. The corrected project data.yaml points the training, validation, and test entries to the folders as extracted on this drive. This matters because a configuration can be syntactically correct while still pointing the training software at the wrong directory.')
table(['Technology or format','Use in this project'],[
('YOLO annotation format','Class IDs and normalized object coordinates in text label files.'),
('YOLOv8 / Ultralytics','Selected target framework for a future detection baseline; no training was run.'),
('Roboflow Universe','Dataset source identified in the supplied metadata; version 3 and CC BY 4.0 are stated in the archive.'),
('Python and Pillow','Used to inspect image readability, dimensions, label pairing, and duplicate hashes.'),
('JSON and YAML','JSON stores reproducible audit results; a corrected YAML config points to this extracted dataset layout.')],[1.65,5.3])
para('The working environment had Python 3.12.14 and Pillow available, but did not have the PyTorch or Ultralytics packages. It reported Intel UHD integrated graphics. This meant a local YOLOv8 training run was not available in the current setup. The archive contained no trained model weights.')

h('5. Internship Activities')
para('The activities recorded here are the concrete project tasks completed during preparation of this report. They should be adjusted to match the student’s actual internship log and supervisor-approved scope.')
for x in ['Extracted the supplied archive into a project folder while preserving its original image and label files.','Read the dataset metadata and recorded its stated source, version, license, export date, class names, resize operation, and augmentation note.','Wrote and ran a Python audit that checked image readability, image-label pairing, YOLO box and polygon row structure, image dimensions, and exact duplicate files across splits.','Created a corrected dataset YAML for the extracted folder layout and saved the results in a machine-readable JSON file.','Reviewed a visual sample from train, validation, and test splits to check image variety and identify annotation questions for follow-up.','Drafted a research protocol describing a baseline training run, validation-based model selection, held-out testing, error analysis, and responsible deployment limits.']: bullet(x)
para('No training or inference activity is listed as completed because the required libraries and trained weights were not available in the supplied environment. The next phase requires a configured PyTorch/Ultralytics runtime and a review of the mixed-format annotation.')
para('The work was organized as a small reproducible pipeline: preserve the source, inventory the files, inspect metadata, audit the labels, review a sample visually, and then write a separate plan for model experiments. Keeping the original data unchanged makes it possible to compare future cleaned or converted copies against the export that was actually received.')

h('6. Technical Learning')
para('A detection dataset is more than a folder of images. Every image needs a matching label file, the class order must stay fixed, and the coordinates must match the task. Reviewing the labels showed why it is useful to inspect actual rows rather than rely only on a README: the training split contains one polygon row among otherwise box-style annotations. A polygon may be useful for segmentation, but this project is framed as object detection, so that annotation should be reviewed and harmonized before a formal training run.')
para('The audit also clarified the difference between dataset quality checks and model evaluation. Counts, image opening, label pairing, and duplicate hashes can show whether files are structurally prepared. They cannot show whether a model detects helmets correctly. For that, a trained model must make predictions on data not used to fit its weights, and its errors must be compared with reliable ground truth.')
para('The supplied split contains 636 training images, 179 validation images, and 88 test images. These are modest sample sizes for a safety-related task, particularly when scenes, cameras, and rider types may vary. Exact duplicate images were not found, but related frames, near-duplicates, and source leakage remain possible because filenames and image hashes do not reveal whether images came from the same original video or capture session.')
para('The dataset labels describe two categories, but the supplied documentation does not define the precise object boundary: the box might mark the helmet/head region or a broader rider region. This is a key technical definition. A consistent class name with inconsistent box meaning can still produce unreliable training targets.')
para('A detector is usually assessed with complementary measures rather than a single accuracy score. Precision describes how many predicted objects are correct; recall describes how many labelled objects were found. Average precision summarizes the precision-recall trade-off for a class, while mAP averages across classes and, depending on the reported form, one or more intersection-over-union thresholds. These measures only become meaningful after a model is trained and evaluated against trustworthy labels on a reserved split.')
para('The class counts are not identical: the audit counted 1,016 helmet annotations and 686 no_helmet annotations. This is a moderate difference in object counts, but it does not reveal whether images, scenes, or difficulty levels are balanced. For that reason, a future evaluation should report results per class and inspect examples, rather than relying only on an overall average.')

h('7. Problems Observed')
table(['Observation','Evidence from the audit','Why it matters'],[
('One mixed polygon annotation','1 polygon row in train; 1,701 other annotation rows use bounding boxes.','The target task is detection. Review and convert the polygon to a box in a derived copy, or make the entire task segmentation.'),
('Class and box policy not described','Metadata supplies class names but not an explicit object-boundary policy.','Annotators may label different objects or regions under the same class.'),
('Small test split','88 test images and 163 annotation rows.','A small test set gives limited evidence about varied road conditions and may produce unstable metrics.'),
('Potential source leakage not assessed','No exact duplicate files found; near-duplicates and related video frames are not ruled out.','Leakage can make validation and test performance look better than performance on new scenes.'),
('Local training runtime unavailable','No PyTorch/Ultralytics installed; no checkpoint in archive.','Training and prediction metrics could not be produced in this environment.'),
('Limited provenance detail','Archive names a Roboflow source and license, but no source-level image record or capture context.','Additional checks are needed for attribution, representativeness, privacy, and external validity.')],[1.35,2.55,3.05])
para('The audit found no missing label files, extra label files, empty labels, unreadable images, malformed box or polygon rows, or exact image duplicates. That is a positive structural result, but it does not establish that the human annotations are correct or that the splits are independent.')
para('The visual sample also showed varied rider types and scene styles, including both motorcycles and bicycles, as well as visible watermarks in some images. This raises a scope question for the research: should the model treat all bicyclists as part of the target population, or is the intended use limited to motorcyclists? Watermarks, stock-like images, or highly polished imagery may also differ from the camera footage expected in a deployment setting. The sample review raises these questions; it does not establish their frequency across the full dataset.')

h('8. Proposed Solutions')
for x in ['Write a short annotation guide defining whether each box encloses a helmet/head or a whole rider, how to label partial occlusion, and what to do when helmet status cannot be judged. Have a second reviewer inspect a sample and resolve disagreements.','In a working copy, visually review the single polygon row. For the detection baseline, convert its polygon to the tight enclosing box and keep the original archive unchanged. Record this transformation in the run log.','Check near-duplicate images and group related frames by original source, capture session, or video before creating splits. Keep groups together so the test set represents genuinely unseen sources.','Add locally representative images across lighting, weather, camera angles, distances, crowding, motorcycle and bicycle scenes, and helmet types. Record these conditions so errors can be analyzed by subgroup.','Train a small YOLOv8 model as a baseline. Use training data to fit weights, validation data to choose checkpoints and confidence thresholds, and the test split only once the model choices are fixed.','Report precision, recall, mAP@0.50 and mAP@0.50:0.95 per class, alongside a confusion matrix, precision-recall curves, and reviewed false positives and false negatives. Retain the environment, configuration, checkpoint, and logs.','Treat model output as a screening aid. Before operational use, test on the target camera and hardware, measure latency, check performance across conditions, and provide human review before any consequential action.']: bullet(x)
para('The project folder includes a reproducible audit script and a research protocol with an example baseline command. That command is proposed guidance only; it was not executed and does not imply any model result.')
table(['Phase','Work product','Decision before moving on'],[
('1. Prepare data','Reviewed labels, agreed class and box policy, harmonized the polygon in a derived copy.','Do the annotations consistently represent the intended objects?'),
('2. Train baseline','Saved configuration, seed, package versions, checkpoint, and training logs.','Can the run be reproduced from the saved files?'),
('3. Evaluate','Per-class metrics, confusion matrix, PR curves, and error examples on held-out data.','Are errors acceptable across important scene conditions?'),
('4. Assess use','Tested target camera, hardware, latency, and human review workflow.','Is there enough evidence for a limited pilot?')],[1.15,3.0,2.8])

h('9. Reflection on Professional Skills')
para('This project reinforced the habit of checking source material before making claims. I initially had a dataset export and a deadline, but the useful work was not to claim a detector result that had not been measured. I inspected the data, recorded what was actually present, and separated those findings from the experiments that still need to be done.')
para('I also practiced organizing a small technical project so another person can follow it: keeping the source dataset intact, writing an audit script, saving a machine-readable result, fixing the path configuration for the extracted folder, and documenting the next steps. The mixed annotation row was a reminder that careful review can uncover issues hidden by a dataset’s broad label description.')
para('For the next stage, I would improve my time planning by checking compute requirements early and reserving time for label review, training, and error analysis. I would also keep a more detailed experiment log from the beginning, including data versions, package versions, configuration changes, and decisions made with the supervisor.')
para('The work also made clear that technical communication is part of responsible development. A useful handover should let a supervisor distinguish what was measured from what is only planned. Recording the consent and source attribution, flagging the limits of the sample, and explaining why no accuracy figure is available are all part of making the project understandable to someone who did not inspect the files personally.')

h('10. Conclusion')
para('The supplied dataset is structurally organized and contains 903 images across three splits, with matching label files and a consistent 640 by 640 image size. The audit counted 1,701 bounding-box annotations and one polygon annotation, with 1,016 class-0 helmet annotations and 686 class-1 no_helmet annotations. No exact duplicate images or basic file-pairing problems were found.')
para('The dataset is not yet sufficient evidence of detector performance. The annotation meaning and polygon row need review; source-level independence and representativeness need further checks; and no training or inference was run. The next defensible milestone is to harmonize annotations, establish a source-grouped evaluation design, train a small baseline in a suitable runtime, and report held-out metrics with examples of errors. This project phase provides a transparent starting point for that work without overstating what has been achieved.')
para('If the next experiment produces promising numbers, the result should still be described as performance on this dataset and its test split. A separate test from the intended region, camera, and user population would be needed to understand whether the result transfers. The project’s most valuable outcome at this stage is a clear, reproducible path from a supplied archive to an experiment that can be evaluated honestly.')

h('11. References')
refs=[('Roboflow Universe. Bike riders helmet detection, version 3. Dataset page and project metadata.','https://universe.roboflow.com/objects-workspace-o5fhb/bike-riders-helmet-detection-bjslq/dataset/3'),('Ultralytics. YOLO object detection dataset format and training guidance.','https://docs.ultralytics.com/datasets/detect/'),('Ultralytics. YOLOv8 model documentation.','https://docs.ultralytics.com/models/yolov8/'),('Ultralytics. YOLOv8 repository and quickstart.','https://github.com/ultralytics/ultralytics'),('Creative Commons. Attribution 4.0 International (CC BY 4.0) legal code.','https://creativecommons.org/licenses/by/4.0/legalcode'),('Supplied archive metadata. README.dataset.txt, README.roboflow.txt, and data.yaml.','Included in the project dataset folder; accessed 25 October 2026.')]
for name,url in refs:
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(2); p.add_run(name)
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.28); p.paragraph_format.space_after=Pt(5)
    r=p.add_run(url); r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(8,126,139)
para('Access date for web references: 25 October 2026.')

h('12. Appendices')
h('Appendix A Dataset audit summary',2)
table(['Split','Images','Label files','Boxes','Polygons','Class 0 helmet','Class 1 no_helmet'],[('Train','636','636','1,240','1','725','516'),('Validation','179','179','298','0','186','112'),('Test','88','88','163','0','105','58'),('Total','903','903','1,701','1','1,016','686')],[.9,.85,1.0,.8,.75,1.25,1.4])
para('Audit method: opened every image with Pillow, checked image/label pairing and extra labels, checked normalized box and polygon row structure, counted dimensions, and computed SHA-256 hashes to identify exact duplicates across splits. Results are saved in analysis/dataset_audit.json. The audit does not verify semantic annotation correctness, near-duplicates, source independence, or model quality.')
h('Appendix B Project files and reproduction',2)
para('Project files include the extracted dataset, audit_dataset.py, analysis/dataset_audit.json, the corrected data.yaml, and research_protocol.md.')
para('The audit was run with Python 3.12.14 and Pillow available in the provided environment. The environment lacked PyTorch and Ultralytics; therefore model training, validation, and prediction were not run.')
h('Appendix C Items to complete before submission',2)
for x in ['Replace bracketed personal, company, supervisor, program, and internship-date fields with verified information; have the host review its profile and reporting structure.','Confirm that the project activities match your internship log and supervisor expectations, then add any institution-required declaration, acknowledgment, signatures, or format.','If you later train a model, add the real environment and measured results; verify dataset attribution and license requirements before submission or public release.']: bullet(x)
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'): p.paragraph_format.keep_with_next=True
props=doc.core_properties; props.title='Motorcycle Rider Helmet Detection Internship Project Report'
props.subject='Dataset assessment and YOLOv8 research development plan'; props.author='[Student name]'
props.keywords='internship, helmet detection, YOLOv8, dataset audit'
doc.save(out)
print(out)
